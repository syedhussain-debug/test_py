import unittest
import time
import random
import requests
import jwt
from appium import webdriver as appium_webdriver
from selenium import webdriver as chorme_webdriver
from appium.options.ios import XCUITestOptions
from appium.webdriver.common.appiumby import AppiumBy
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from faker import Faker
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.chrome.service import Service
from webdriver_manager.chrome import ChromeDriverManager
from getotp import OTPService
from openpyxl import Workbook, load_workbook
from logout import LogoutFunction
from setupFile import SetupFile

import os

EXCEL_FILE = "phone_numbers.xlsx"
SHEET_NAME = "Numbers"

class TestSignupWithOTP(SetupFile):
    

        

    def test_signup(self):
        self.fake = Faker("ar_SA")
        """Verify user can sign up using OTP"""
        logout = LogoutFunction()
        logout.driver = self.driver
        logout.wait = self.wait
        logout.logoutFunction()
        # 📱 Phone number
        contact_no = self.wait.until(
            EC.visibility_of_element_located((By.NAME, "XXXXXXXXX"))
        )
        phone_number = self.generate_random_number()
        contact_no.clear()
        contact_no.send_keys(phone_number)
        print(f"📱 Using phone number: {phone_number}")

        # 👤 First name
        first_name = self.generate_random_name()
        first_name_field = self.wait.until(
            EC.visibility_of_element_located((By.NAME, "Your name"))
        )
        first_name_field.send_keys(first_name)

        # 👤 Last name
        last_name = self.generate_random_last_name()
        last_name_field = self.wait.until(
            EC.visibility_of_element_located((By.NAME, "Your last name"))
        )
        last_name_field.send_keys(last_name)

        # 📧 Email
        email = self.generate_random_email()
        email_field = self.wait.until(
            EC.visibility_of_element_located((By.NAME, "example@company.com"))
        )
        email_field.send_keys(email)
        self.wait.until(EC.element_to_be_clickable((By.ID, "Last Name"))).click()
        self.save_user_data(phone_number, first_name, last_name, email)


        # ➡️ Next
        self.wait.until(EC.element_to_be_clickable((By.ID, "Next"))).click()

        # 🔐 Fetch OTP from API
        #otp_code = "1234"
        get_otp_service = OTPService()
        otp_code = get_otp_service.get_otp_code(phone_number)
        print(f"🔐 OTP Received: {otp_code}")

        # ✍️ Enter OTP
        otp_field = self.wait.until(
            EC.visibility_of_element_located(
                (AppiumBy.CLASS_NAME, "XCUIElementTypeTextField")
            )
        )
        otp_field.send_keys(otp_code)

    # ---------------- HELPERS ---------------- #

    def generate_random_number(self):
        prefixes = ['50', '53', '54', '55', '56', '57', '58', '59']
        prefix = random.choice(prefixes)
        rest = ''.join(str(random.randint(0, 9)) for _ in range(7))
        return f"{prefix}{rest}"

    def generate_random_name(self):
        return self.fake.first_name()

    def generate_random_last_name(self):
        return self.fake.last_name()

    def generate_random_email(self):
        return self.fake.email()

    def save_user_data(self, phone, first_name, last_name, email):
    # Create file if not exists
        if not os.path.exists(EXCEL_FILE):
            wb = Workbook()
            ws = wb.active
            ws.title = SHEET_NAME
            ws.append(["Phone Number", "First Name", "Last Name", "Email"])  # Header row
            wb.save(EXCEL_FILE)

    # Load workbook
        wb = load_workbook(EXCEL_FILE)

        # Create sheet if missing
        if SHEET_NAME not in wb.sheetnames:
            ws = wb.create_sheet(SHEET_NAME)
            ws.append(["Phone Number", "First Name", "Last Name", "Email"])
        else:
            ws = wb[SHEET_NAME]

        # Append user data
        ws.append([phone, first_name, last_name, email])
        wb.save(EXCEL_FILE)

        print(f"✅ User data saved: {phone}, {first_name} {last_name}, {email}")


    def tearDown(self):
        if self.driver:
            self.driver.quit()




if __name__ == "__main__":
    unittest.main()
