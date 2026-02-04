import unittest
import sys
import os
current_dir = os.path.dirname(os.path.abspath(__file__))
parent_dir = os.path.dirname(current_dir)
sys.path.insert(0, parent_dir)
from selenium.webdriver.common.by import By
from selenium.webdriver.support import expected_conditions as EC
from appium.webdriver.common.appiumby import AppiumBy
from functions.logout import LogoutFunction
from functions.getotp import OTPService
from openpyxl import load_workbook

from test.initializeByAlert import InitializeByAlert
from test.profileStatus import ProfileStatus
import random
from functions.setupFile import SetupFile


class LoginFunction(SetupFile):

    def test_login(self):
        print("🚀 Starting login test...")
        """Verify user can log in"""
        # 🔁 Step 1: Check login state
        logout = LogoutFunction()
        logout.driver = self.driver
        logout.wait = self.wait
        logout.logoutFunction()

        # Additional login steps would go here
        contact_no = self.wait.until(
            EC.visibility_of_element_located((By.NAME, "XXXXXXXXX"))
        )
        phone_number = self.get_random_phone_from_excel()
        contact_no.clear()
        contact_no.send_keys(phone_number)
        print(f"📱 Using phone number: {phone_number}")

        get_otp_service = OTPService()
        otp_code = get_otp_service.get_otp_code(phone_number)
        print(f"🔐 OTP Received: {otp_code}")

        # ✍️ Enter OTP
        otp_field = self.wait.until(
            EC.visibility_of_element_located(
                (AppiumBy.CLASS_NAME, "XCUIElementTypeTextField")
            )
        )
        otp_field.send_keys(otp_code)

        locationAndAlert = InitializeByAlert()
        locationAndAlert.intilize_setup(self.driver, self.wait)        
        profileStatus = ProfileStatus()
        profileStatus.assertionCondition(self.driver, self.wait, self)

        
    def get_random_phone_from_excel(self):
        file_path = "phone_numbers.xlsx"
        sheet_name = "Numbers"

        wb = load_workbook(file_path)
        ws = wb[sheet_name]

        phone_numbers = []

        # Skip header row (row 1)
        for row in ws.iter_rows(min_row=2, max_col=1, values_only=True):
            if row[0]:  # make sure cell not empty
                phone_numbers.append(str(row[0]))

        wb.close()

        if not phone_numbers:
            raise Exception("❌ No phone numbers found in Excel!")

        selected_number = random.choice(phone_numbers)
        print(f"🎲 Random phone number selected: {selected_number}")
        return selected_number


    def tearDown(self):
        if self.driver:
            self.driver.quit()


if __name__ == "__main__":
    unittest.main()
