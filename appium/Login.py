import unittest
from appium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.common.exceptions import TimeoutException
from appium.webdriver.common.appiumby import AppiumBy
from appium.options.ios import XCUITestOptions
from logout import LogoutFunction
from getotp import OTPService
from openpyxl import load_workbook
import random


class LoginFunction(unittest.TestCase):

    def setUp(self):
        options = XCUITestOptions()
        options.platform_name = "iOS"
        options.platform_version = "17.0"
        options.device_name = "iPhone 11"
        options.udid = "00008030-000621290C39802E"
        options.automation_name = "XCUITest"
        options.bundle_id = "com.yumealz.uat"
        options.no_reset = True
        options.show_xcode_log = True
        options.wda_startup_retries = 1

        self.driver = webdriver.Remote(
            command_executor="http://127.0.0.1:4723",
            options=options
        )
        self.wait = WebDriverWait(self.driver, 30)
    def test_login(self):
        """Verify user can log in"""

        # 🔁 Step 1: Check login state
        logout = LogoutFunction()
        logout.driver = self.driver
        logout.wait = self.wait
        logout.logoutFunction()

        # Additional login steps would go here
        contact_no = self.wait.until(
            EC.visibility_of_element_located((By.NAME, "XXXXXXXXX"))
        )
        phone_number = self.get_random_phone_from_excel()
        contact_no.clear()
        contact_no.send_keys(phone_number)
        print(f"📱 Using phone number: {phone_number}")

        get_otp_service = OTPService()
        otp_code = get_otp_service.get_otp_code(phone_number)
        print(f"🔐 OTP Received: {otp_code}")

        # ✍️ Enter OTP
        otp_field = self.wait.until(
            EC.visibility_of_element_located(
                (AppiumBy.CLASS_NAME, "XCUIElementTypeTextField")
            )
        )
        otp_field.send_keys(otp_code)
    def get_random_phone_from_excel(self):
        file_path = "phone_numbers.xlsx"
        sheet_name = "Numbers"

        wb = load_workbook(file_path)
        ws = wb[sheet_name]

        phone_numbers = []

        # Skip header row (row 1)
        for row in ws.iter_rows(min_row=2, max_col=1, values_only=True):
            if row[0]:  # make sure cell not empty
                phone_numbers.append(str(row[0]))

        wb.close()

        if not phone_numbers:
            raise Exception("❌ No phone numbers found in Excel!")

        selected_number = random.choice(phone_numbers)
        print(f"🎲 Random phone number selected: {selected_number}")
        return selected_number


    def tearDown(self):
        if self.driver:
            self.driver.quit()


if __name__ == "__main__":
    unittest.main()
